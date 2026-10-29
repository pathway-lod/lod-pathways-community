#!/usr/bin/env python3
"""
Pilot: ipecac alkaloid biosynthesis, part 2: the paired data (needs the supplementary files of the paper, see the README).

Uses the Supplementary Information of Colinas et al. (article 41589_2025_1926), placed in supplementaries/:
  MOESM3  protein sequences of the cloned genes and of the reference genes (OMT, glucosidase and reductase trees)
  MOESM4  Source Data Figure 2: peak areas of the metabolites per tissue (3 plants), substrate-feeding experiments
  MOESM5-7 Source Data Figures 3-5: heterologous expression in N. benthamiana

Outputs in results/:
  tissue_pairing.csv                         which transcriptome column is paired with which metabolite tissue, with the check
  gene_metabolite_summary.csv                transcripts correlated with each metabolite across the tissues, and what chance would give
  gene_metabolite_top_{Ci,As}.csv            the best-correlated transcripts per metabolite
  gene_metabolite_class_summary.csv          what the correlated transcripts are (family level)
  sequence_identity_to_plantmetwiki.csv      each cloned gene against the enzymes of PlantMetWiki's emetine pathway
  reaction_evidence.csv                      what each infiltration or heterologous expression experiment shows, against its control

  python pair_ipecac.py
"""

import glob
import re
import sys
import warnings

import numpy as np
import pandas as pd
import requests
from Bio import Align
from Bio.Align import substitution_matrices
from scipy import stats

warnings.filterwarnings('ignore')
import match_ipecac as base
from match_ipecac import CLASSES, DATA, PEARSON_MIN, RESULTS, SETS, class_summary, read_table

HERE = base.HERE
SUPP = HERE / "supplementaries"
EXPR_MIN = base.EXPR_MIN

# transcriptome column -> tissue label of the metabolite tables. The six columns of Carapichea and the twelve of Alangium
# are given by the paper's tissue list (Supplementary Fig. 1); the abbreviations of the transcriptome files are not
# explained in the deposit, so the assignment is a reading and is tested (see tissue_check).
TISSUES = {
    "Ci": {"YL": "young leaves", "ML": "expanded leaves", "OL": "old leaves", "Stem": "stem", "MR": "rhizome", "LR": "roots"},
    "As": {"As_LB": "leaf buds", "As_L1": "leaf 1", "As_L2": "leaf 2", "As_L3": "leaf 3", "As_L4": "leaf 4", "As_L5": "leaf 5",
           "As_L6": "leaf 6", "As_L7": "leaf 7", "As_GS": "young stem", "As_WS": "old stem", "As_SB": "stem bark", "As_R": "roots"},
}
FIG2 = {"Ci": "Fig. 2e", "As": "Fig. 2f"}
UNIPROT_ENZYMES = ["D3KY98", "D3KY99", "D3KYA0", "B6ZKM3", "E7FL15"]    # enzymes of PWY-7135 in PlantMetWiki (see match_ipecac.py)


def norm(s):
    return " ".join(str(s).split())          # the source tables contain non-breaking spaces


def supp(n):
    hits = sorted(glob.glob(str(SUPP / f"*MOESM{n}_ESM*.xlsx")))
    if not hits:
        sys.exit(f"supplementaries/ has no *MOESM{n}_ESM*.xlsx; download the Supplementary Information of the paper (see the README)")
    return hits[0]


# ------------------------------------------------------------------------------------------------- source data
def sheet_rows(path, sheet):
    import openpyxl
    return list(openpyxl.load_workbook(path, data_only=True)[sheet].iter_rows(values_only=True))


def data_rows(rows):
    """Rows of a Source Data sheet that hold numbers: a label in the first cell, a number in the second."""
    return [r for r in rows[4:] if r[0] and isinstance(r[1], (int, float))]


def groups(header, rows):
    """Column group (compound or condition) of every numeric column; the sheets carry empty columns to the right."""
    drows = data_rows(rows)
    width = max(max((i for i, x in enumerate(r[1:]) if x is not None), default=-1) for r in drows) + 1
    names, cur = [], None
    for c in header[1:1 + width]:
        if c:
            cur = norm(c)
        names.append(cur)
    return names


def tissue_matrix(species):
    """metabolite x tissue means of the three plants, and the replicate values."""
    rows = sheet_rows(supp(4), FIG2[species])
    names = groups(rows[3], rows)
    reps = {}
    for r in data_rows(rows):
        v = np.array([np.nan if x is None else float(x) for x in r[1:1 + len(names)]])
        for m in dict.fromkeys(names):
            reps.setdefault(m, {})[norm(r[0])] = v[[i for i, n in enumerate(names) if n == m]]
    mean = pd.DataFrame({m: {t: x.mean() for t, x in d.items()} for m, d in reps.items()})
    return mean, reps


def evidence_long():
    """All infiltration and heterologous expression experiments as (sheet, condition, compound, replicates, control)."""
    out = []
    for fig, n, sheet, layout in [("2g", 4, "Fig. 2g", "time"), ("2h", 4, "Fig. 2h", "time"), ("2i", 4, "Fig. 2i", "time"),
                                  ("3b", 5, "Fig. 3b", "genes"), ("3c", 5, "Fig. 3c", "genes"), ("4c", 6, "Fig. 4c", "genes"),
                                  ("4f", 6, "Fig. 4f", "genes"), ("5b", 7, "Fig. 5b", "genes"), ("5d", 7, "Fig. 5d", "genes")]:
        rows = sheet_rows(supp(n), sheet)
        names = groups(rows[3], rows)
        title = norm(rows[1][0])
        for r in data_rows(rows):
            label = norm(r[0])
            v = np.array([np.nan if x is None else float(x) for x in r[1:1 + len(names)]])
            for g in dict.fromkeys(names):
                reps = v[[i for i, nm in enumerate(names) if nm == g]]
                if layout == "time":     # rows are compounds, column groups are conditions
                    out.append({"figure": fig, "title": title, "condition": g, "compound": label, "replicates": reps,
                                "control": "0h" if g == "24h" else ("Mock" if g.startswith("1 mM") else None)})
                else:                    # rows are the genes expressed, column groups are compounds
                    out.append({"figure": fig, "title": title, "condition": label, "compound": g, "replicates": reps,
                                "control": "EV" if label != "EV" else None})
    return out


def reaction_evidence(matches):
    ev = evidence_long()
    by = {(e["figure"], e["condition"], e["compound"]): e["replicates"] for e in ev}
    paper = dict(zip(matches.name, zip(matches.match_level, matches.plantmetwiki_label)))
    rows = []
    for e in ev:
        if e["control"] is None:
            continue
        ctrl = by[(e["figure"], e["control"], e["compound"])]
        m, c = e["replicates"].mean(), ctrl.mean()
        with np.errstate(all="ignore"):
            p = stats.ttest_ind(e["replicates"], ctrl, equal_var=False).pvalue if np.std(e["replicates"]) + np.std(ctrl) > 0 else np.nan
        # with three plants per condition a t-test has little power; the call asks for a 3-fold difference of the means
        # and complete separation of the replicates (every value of one group beyond every value of the other)
        r, k = e["replicates"], ctrl
        if m >= 3 * c and m > 0 and r.min() > k.max():
            call = "higher"
        elif c >= 3 * m and c > 0 and k.min() > r.max():
            call = "lower"
        else:
            call = "no clear change"
        name = re.sub(r"\s+\d+[a-z]?$", "", e["compound"])      # "6-O-Me-DAIIA 6a" -> "6-O-Me-DAIIA"
        lvl, lab = paper.get(name, ("no structure in the deposit", ""))
        rows.append({"figure": e["figure"], "experiment": e["title"].replace("Normalised peak areas of three biological replicates of ", ""),
                     "condition": e["condition"], "control": e["control"], "compound": e["compound"],
                     "mean": round(m, 1), "sd": round(float(np.std(e["replicates"], ddof=1)), 1), "control_mean": round(c, 1),
                     "fold_vs_control": round(m / c, 1) if c > 0 else np.nan, "p_welch_n3": p, "call": call,
                     "compound_vs_plantmetwiki": lvl, "plantmetwiki_label": lab})
    return pd.DataFrame(rows)


# ------------------------------------------------------------------------------------------------- pairing
def zscores(x):
    x = np.asarray(x, float)
    sd = x.std(axis=-1, keepdims=True)
    with np.errstate(all="ignore"):
        return (x - x.mean(axis=-1, keepdims=True)) / sd


def pairing(species, matches):
    file, desc, _ = SETS[species]
    d = read_table(file)
    cols = list(TISSUES[species])
    mean, _ = tissue_matrix(species)
    k = len(cols)
    X = d[cols].astype(float).values                         # transcripts x tissues
    ZX = zscores(X)
    metab = {m: np.array([mean.loc[TISSUES[species][c], m] for c in cols]) for m in mean.columns}
    bait = d[np.isclose(d.Pearson, 1.0)].iloc[0]
    zb = zscores(bait[cols].astype(float).values[None, :])[0]
    check = {m: float(np.mean(zb * zscores(v[None, :])[0])) for m, v in metab.items()}   # r with n tissues, population z-scores
    # the same check with the two tissues of each pair swapped, for the reading of the abbreviations
    swap = {"Ci": ("MR", "LR"), "As": ("As_GS", "As_WS")}[species]
    alt = dict(TISSUES[species]); alt[swap[0]], alt[swap[1]] = alt[swap[1]], alt[swap[0]]
    check_alt = {m: float(np.corrcoef(bait[cols].astype(float).values, [mean.loc[alt[c], m] for c in cols])[0, 1]) for m in mean.columns}
    p_null = stats.t.sf(PEARSON_MIN * np.sqrt((k - 2) / (1 - PEARSON_MIN ** 2)), k - 2)
    pmax = d[cols].max(axis=1).values
    n_expr = int(((pmax >= EXPR_MIN) & d[cols].notna().all(axis=1).values & (X.std(axis=1) > 0)).sum())
    summary, tops = [], []
    for m, v in metab.items():
        r = (ZX * zscores(v[None, :])[0]).mean(axis=1)       # Pearson r of every transcript with the metabolite
        ok = ~np.isnan(r) & (pmax >= EXPR_MIN)
        n75 = int(((r >= PEARSON_MIN) & ok).sum())
        summary.append({"species": species, "metabolite": m, "tissues": k, "bait_r_with_metabolite": round(check[m], 2),
                        "bait_r_if_tissues_swapped": round(check_alt[m], 2), "transcripts_r_ge_0.75_and_expr_ge_50": n75,
                        "transcripts_expr_ge_50": n_expr, "expected_by_chance_r_ge_0.75": round(n_expr * p_null), "p_at_r_0.75_two_sided": round(2 * p_null, 4)})
        idx = np.argsort(-np.where(ok, r, -2))[:100]
        top = d.iloc[idx][["Name", desc]].rename(columns={desc: "description"})
        top["metabolite"], top["r"], top["max_expression"] = m, r[idx], pmax[idx]
        tops.append(top)
    # class enrichment among transcripts correlated with at least one metabolite
    rall = np.vstack([(ZX * zscores(v[None, :])[0]).mean(axis=1) for v in metab.values()])
    chosen = d[((np.nanmax(np.where(np.isnan(rall), -2, rall), axis=0) >= PEARSON_MIN) & (pmax >= EXPR_MIN))].copy()
    chosen = chosen.rename(columns={desc: "description"})
    chosen["description"] = chosen.description.fillna("none")
    cls = class_summary(species, d, desc, chosen)
    pairing_rows = [{"species": species, "transcriptome_column": c, "metabolite_tissue": TISSUES[species][c]} for c in cols]
    return pairing_rows, summary, pd.concat(tops), cls, len(chosen)


# ------------------------------------------------------------------------------------------------- sequences
def read_fasta(text):
    seqs, cur = {}, None
    for line in text.splitlines():
        if line.startswith(">"):
            cur = line[1:].split("|")[1] if "|" in line else line[1:].split()[0]
            seqs[cur] = ""
        elif cur:
            seqs[cur] += line.strip()
    return seqs


def uniprot_sequences():
    path = DATA / "uniprot_pwy7135_sequences.fasta"
    if not path.exists():
        r = requests.get("https://rest.uniprot.org/uniprotkb/accessions", params={"accessions": ",".join(UNIPROT_ENZYMES), "format": "fasta"}, timeout=120)
        r.raise_for_status()
        path.write_text(r.text)
    return read_fasta(path.read_text())


def study_proteins():
    import openpyxl
    wb = openpyxl.load_workbook(supp(3), data_only=True)
    out = []
    for ws in wb.worksheets:
        for r in ws.iter_rows(min_row=2, values_only=True):
            if r[2] and "this study" in str(r[2]).lower() and r[1] and r[4]:
                out.append({"tree": ws.title, "gene": norm(r[1]), "species": norm(r[0]), "sequence": "".join(str(r[4]).split())})
    return out


FAMILY_REFERENCE = {"OMT tree": ["D3KY98", "D3KY99", "D3KYA0", "E7FL15"],      # the O-methyltransferases of PWY-7135
                    "Glucosidase tree": ["B6ZKM3"],                           # IpeGLU1
                    "Reductase tree": []}                                     # no reductase in PWY-7135


def identity(a, b):
    al = Align.PairwiseAligner()
    al.substitution_matrix = substitution_matrices.load("BLOSUM62")
    al.open_gap_score, al.extend_gap_score = -10, -0.5
    al.target_end_gap_score = al.query_end_gap_score = 0       # end gaps are free, so a partial sequence is not punished
    aln = al.align(a, b)[0]
    return 100 * aln.counts().identities / min(len(a), len(b))


def sequence_table():
    ref = uniprot_sequences()
    info = pd.read_csv(DATA / "uniprot_pwy7135_enzymes.tsv", sep="\t").set_index("Entry")
    rows = []
    for g in study_proteins():
        scored = sorted(((identity(g["sequence"], ref[acc]), acc) for acc in FAMILY_REFERENCE[g["tree"]]), reverse=True)
        if not scored:
            rows.append({"gene": g["gene"], "species": g["species"], "tree": g["tree"], "length": len(g["sequence"]),
                         "best_plantmetwiki_enzyme": "", "best_enzyme_name": "no enzyme of this family in PWY-7135", "identity_percent": np.nan, "all_identities": ""})
            continue
        best = scored[0]
        rows.append({"gene": g["gene"], "species": g["species"], "tree": g["tree"], "length": len(g["sequence"]),
                     "best_plantmetwiki_enzyme": best[1], "best_enzyme_name": info.loc[best[1], "Protein names"].split(" (")[0],
                     "identity_percent": round(best[0], 1),
                     "all_identities": "; ".join(f"{a} {i:.0f}" for i, a in scored)})
    return pd.DataFrame(rows)


def main():
    RESULTS.mkdir(exist_ok=True)
    matches = pd.read_csv(RESULTS / "metabolite_matches.csv")
    pair_rows, summ, tops, cls = [], [], {}, []
    for sp in ("Ci", "As"):
        pr, s, t, c, n = pairing(sp, matches)
        pair_rows += pr
        summ += s
        tops[sp] = t
        cls += c
        print(f"{sp}: {n} transcripts have r >= {PEARSON_MIN} with at least one metabolite and expression >= {EXPR_MIN}")
    pd.DataFrame(pair_rows).to_csv(RESULTS / "tissue_pairing.csv", index=False)
    S = pd.DataFrame(summ)
    S.to_csv(RESULTS / "gene_metabolite_summary.csv", index=False)
    for sp, t in tops.items():
        t.to_csv(RESULTS / f"gene_metabolite_top_{sp}.csv", index=False)
    print("\n== bait gene against the metabolites over the tissues (tests the tissue pairing)")
    print(S[["species", "metabolite", "bait_r_with_metabolite", "bait_r_if_tissues_swapped"]].to_string(index=False))
    print("\n== transcripts correlated with each metabolite")
    print(S[["species", "metabolite", "tissues", "transcripts_expr_ge_50", "transcripts_r_ge_0.75_and_expr_ge_50", "expected_by_chance_r_ge_0.75"]].to_string(index=False))
    cdf = pd.DataFrame(cls)
    p = cdf.fisher_p_greater.values
    order = np.argsort(p)
    adj = np.empty_like(p)
    adj[order] = np.minimum.accumulate((p[order] * len(p) / (np.arange(len(p)) + 1))[::-1])[::-1].clip(max=1)
    cdf["fdr_bh"] = adj
    cdf.to_csv(RESULTS / "gene_metabolite_class_summary.csv", index=False)
    print("\n== classes among transcripts correlated with a metabolite (FDR < 0.1 only)")
    print(cdf[cdf.fdr_bh < 0.1][["transcriptome", "class", "co_expressed", "in_all", "fold_enrichment", "fdr_bh"]].round(4).to_string(index=False))

    seq = sequence_table()
    seq.to_csv(RESULTS / "sequence_identity_to_plantmetwiki.csv", index=False)
    print("\n== cloned genes against the enzymes of PlantMetWiki's emetine pathway")
    print(seq[["gene", "species", "best_plantmetwiki_enzyme", "best_enzyme_name", "identity_percent"]].to_string(index=False))

    ev = reaction_evidence(matches)
    ev.to_csv(RESULTS / "reaction_evidence.csv", index=False)
    print("\n== experiments: increases and decreases against the control")
    print(ev[ev.call != "no clear change"][["figure", "condition", "compound", "mean", "control_mean", "call"]].to_string(index=False, max_colwidth=48))


if __name__ == "__main__":
    main()
